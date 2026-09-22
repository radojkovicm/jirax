"""Excel export helpers for JiraX Time Tracker."""
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment

HEADER_FONT = Font(bold=True, size=12)
HEADER_FILL = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal='center', vertical='center', wrap_text=True)
TOTAL_FILL = PatternFill(start_color="E6F2FF", end_color="E6F2FF", fill_type="solid")


def format_sheet(worksheet, df):
    """Bold/shade the header row and auto-size columns for a DataFrame sheet."""
    for col_num, column in enumerate(df.columns, 1):
        cell = worksheet.cell(row=1, column=col_num)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT

    for idx, col in enumerate(df.columns):
        if not df.empty:
            max_length = max(df[col].astype(str).apply(len).max(), len(col))
        else:
            max_length = len(col)
        worksheet.column_dimensions[chr(65 + idx)].width = max_length + 2


def write_task_report(file_path, rows):
    """Write the 'Reports' tab rows to a single formatted Excel sheet with summaries."""
    columns = ["Datum", "Korisnik", "Zadatak", "Kategorija", "Sati", "Početak", "Kraj", "Prioritet"]
    df = pd.DataFrame(rows, columns=columns)

    with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Zadaci')
        worksheet = writer.sheets['Zadaci']
        format_sheet(worksheet, df)

        total_hours = df["Sati"].sum() if not df.empty else 0.0

        total_row_idx = worksheet.max_row + 2
        worksheet.cell(row=total_row_idx, column=1, value="UKUPNO").font = Font(bold=True)
        worksheet.cell(row=total_row_idx, column=1).fill = TOTAL_FILL

        hours_col_idx = list(df.columns).index("Sati") + 1
        cell = worksheet.cell(row=total_row_idx, column=hours_col_idx, value=total_hours)
        cell.font = Font(bold=True)
        cell.fill = TOTAL_FILL

        next_row = total_row_idx + 2
        if not df.empty:
            next_row = _write_breakdown(worksheet, df, "Kategorija", "Statistika po kategorijama", next_row, total_hours)
            next_row = _write_breakdown(worksheet, df, "Korisnik", "Statistika po korisnicima", next_row + 1, total_hours)


def _write_breakdown(worksheet, df, group_col, title, start_row, total_hours):
    stats = df.groupby(group_col)["Sati"].sum().reset_index().sort_values("Sati", ascending=False)

    worksheet.cell(row=start_row, column=1, value=title).font = Font(bold=True, size=12)
    header_row = start_row + 1
    for col_idx, header_text in enumerate([group_col, "Sati", "Procenat"], 1):
        cell = worksheet.cell(row=header_row, column=col_idx, value=header_text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT

    row = header_row + 1
    for _, data_row in stats.iterrows():
        percentage = (data_row["Sati"] / total_hours * 100) if total_hours > 0 else 0
        worksheet.cell(row=row, column=1, value=data_row[group_col])
        worksheet.cell(row=row, column=2, value=data_row["Sati"])
        worksheet.cell(row=row, column=3, value=f"{percentage:.1f}%")
        row += 1
    return row


def write_statistics_report(file_path, category_stats, user_stats, daily_stats, details_rows):
    """Write the 'Stats' tab export: per-category, per-user, daily and detail sheets."""
    total_hours = sum(hours for _, hours in category_stats)

    cat_data = [[name, hours, (hours / total_hours * 100) if total_hours > 0 else 0]
                for name, hours in category_stats]
    cat_df = pd.DataFrame(cat_data, columns=["Kategorija", "Sati", "Procenat"])

    user_data = [[name or "Nepoznat", hours, (hours / total_hours * 100) if total_hours > 0 else 0]
                 for name, hours in user_stats]
    user_df = pd.DataFrame(user_data, columns=["Korisnik", "Sati", "Procenat"])

    daily_data = [[date, category, user or "Nepoznat", hours] for date, category, user, hours in daily_stats]
    daily_df = pd.DataFrame(daily_data, columns=["Datum", "Kategorija", "Korisnik", "Sati"])

    details_df = pd.DataFrame(
        [list(row) for row in details_rows],
        columns=["Datum", "Korisnik", "Kategorija", "Zadatak", "Sati", "Početak", "Kraj", "Prioritet", "Opis"],
    )

    with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
        cat_df.to_excel(writer, index=False, sheet_name='Po kategorijama')
        user_df.to_excel(writer, index=False, sheet_name='Po korisnicima')
        daily_df.to_excel(writer, index=False, sheet_name='Dnevna statistika')
        details_df.to_excel(writer, index=False, sheet_name='Detalji')

        for sheet_name, df in (
            ('Po kategorijama', cat_df),
            ('Po korisnicima', user_df),
            ('Dnevna statistika', daily_df),
            ('Detalji', details_df),
        ):
            format_sheet(writer.sheets[sheet_name], df)
